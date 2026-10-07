"""ATS Resume Checker - Streamlit + Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description,
and get an ATS score with concrete suggestions for improvement.
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL_NAME = "gemini-3.5-flash"
MAX_RESUME_CHARS = 20000  # keeps the prompt small and fast
MIN_RESUME_CHARS = 200  # below this the file is probably scanned / empty

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. You evaluate how well a resume will parse and rank \
in an ATS, and how strong it is for human recruiters.

Be honest, specific and consistent. Do not inflate scores. Base every comment on \
the actual resume text; never invent experience the candidate does not have.

Return ONLY a JSON object with exactly this structure:
{
  "overall_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "category_scores": {
    "formatting_and_structure": <integer 0-100>,
    "keywords_and_relevance": <integer 0-100>,
    "experience_and_impact": <integer 0-100>,
    "skills": <integer 0-100>,
    "education_and_certifications": <integer 0-100>,
    "clarity_and_grammar": <integer 0-100>
  },
  "strengths": ["<short bullet>", ...],
  "weaknesses": ["<short bullet>", ...],
  "missing_keywords": ["<keyword or skill>", ...],
  "improvements": [
    {
      "priority": "High" | "Medium" | "Low",
      "section": "<resume section, e.g. Summary, Experience, Skills>",
      "issue": "<what is wrong>",
      "suggestion": "<exactly how to fix it>",
      "example": "<a rewritten example line, or empty string>"
    }
  ]
}

Scoring guidance:
- If a job description is provided, weight keyword match against it heavily and \
list the important job-description keywords that are missing from the resume.
- If no job description is provided, judge general ATS readiness and list \
commonly expected keywords for the candidate's apparent target role.
- Reward: standard section headings, consistent dates, quantified achievements, \
strong action verbs, relevant skills, clean single-column text.
- Penalize: missing contact info, missing sections, vague duties, no metrics, \
typos, very long paragraphs, keyword gaps.
- Give 3-6 strengths, 3-6 weaknesses, up to 15 missing keywords and 5-8 \
improvements ordered by priority (High first)."""

CATEGORY_LABELS = {
    "formatting_and_structure": "Formatting & Structure",
    "keywords_and_relevance": "Keywords & Relevance",
    "experience_and_impact": "Experience & Impact",
    "skills": "Skills",
    "education_and_certifications": "Education & Certifications",
    "clarity_and_grammar": "Clarity & Grammar",
}


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many resumes put content in tables (two-column layouts).
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    parts.append(text)
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# --------------------------------------------------------------------------
# Gemini
# --------------------------------------------------------------------------
def get_api_key(sidebar_key: str = "") -> str:
    if sidebar_key.strip():
        return sidebar_key.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return str(st.secrets["GEMINI_API_KEY"]).strip()
    except Exception:
        pass  # no secrets file locally - that's fine
    return os.environ.get("GEMINI_API_KEY", "").strip()


def build_user_prompt(resume_text: str, job_description: str) -> str:
    prompt = f"RESUME:\n\"\"\"\n{resume_text[:MAX_RESUME_CHARS]}\n\"\"\"\n\n"
    if job_description.strip():
        prompt += f"JOB DESCRIPTION:\n\"\"\"\n{job_description.strip()[:8000]}\n\"\"\"\n"
    else:
        prompt += "JOB DESCRIPTION: (none provided - do a general ATS review)\n"
    return prompt


def parse_json_response(raw: str) -> dict:
    """Parse model output as JSON, tolerating code fences and stray text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start : end + 1])
        raise


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Make sure the result has every field the UI needs, with safe types."""
    if not isinstance(data, dict):
        raise ValueError("Unexpected response format from the model.")

    scores = data.get("category_scores") or {}
    category_scores = {k: _clamp(scores.get(k)) for k in CATEGORY_LABELS}

    def str_list(key):
        items = data.get(key) or []
        return [str(i).strip() for i in items if str(i).strip()] if isinstance(items, list) else []

    improvements = []
    for item in data.get("improvements") or []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "Medium")).capitalize()
        if priority not in ("High", "Medium", "Low"):
            priority = "Medium"
        improvements.append(
            {
                "priority": priority,
                "section": str(item.get("section", "General")),
                "issue": str(item.get("issue", "")),
                "suggestion": str(item.get("suggestion", "")),
                "example": str(item.get("example", "") or ""),
            }
        )
    order = {"High": 0, "Medium": 1, "Low": 2}
    improvements.sort(key=lambda i: order[i["priority"]])

    overall = data.get("overall_score")
    if overall is None and category_scores:
        overall = sum(category_scores.values()) / len(category_scores)

    return {
        "overall_score": _clamp(overall),
        "summary": str(data.get("summary", "")),
        "category_scores": category_scores,
        "strengths": str_list("strengths"),
        "weaknesses": str_list("weaknesses"),
        "missing_keywords": str_list("missing_keywords"),
        "improvements": improvements,
    }


def analyze_resume(api_key: str, resume_text: str, job_description: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=build_user_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    if not getattr(response, "text", None):
        raise ValueError("The model returned an empty response. Please try again.")
    return normalize_result(parse_json_response(response.text))


# --------------------------------------------------------------------------
# UI helpers
# --------------------------------------------------------------------------
def score_color(score: int) -> str:
    if score >= 80:
        return "#16a34a"  # green
    if score >= 60:
        return "#d97706"  # amber
    return "#dc2626"  # red


def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent - ATS ready"
    if score >= 60:
        return "Good - needs some work"
    if score >= 40:
        return "Fair - significant improvements needed"
    return "Poor - major revisions needed"


def render_score_card(score: int) -> None:
    color = score_color(score)
    st.markdown(
        f"""
        <div style="text-align:center;padding:1.2rem;border-radius:16px;
                    border:2px solid {color};">
            <div style="font-size:4rem;font-weight:800;color:{color};line-height:1;">
                {score}<span style="font-size:1.5rem;">/100</span>
            </div>
            <div style="font-size:1.05rem;margin-top:.4rem;">{score_label(score)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_results(result: dict) -> None:
    left, right = st.columns([1, 2], gap="large")
    with left:
        render_score_card(result["overall_score"])
    with right:
        st.subheader("Summary")
        st.write(result["summary"] or "No summary returned.")

    st.divider()
    st.subheader("Score breakdown")
    cols = st.columns(3)
    for i, (key, label) in enumerate(CATEGORY_LABELS.items()):
        score = result["category_scores"][key]
        with cols[i % 3]:
            st.metric(label, f"{score}/100")
            st.progress(score / 100)

    st.divider()
    col_a, col_b = st.columns(2, gap="large")
    with col_a:
        st.subheader("Strengths")
        for s in result["strengths"] or ["None identified."]:
            st.markdown(f"- {s}")
    with col_b:
        st.subheader("Weaknesses")
        for w in result["weaknesses"] or ["None identified."]:
            st.markdown(f"- {w}")

    st.divider()
    st.subheader("Missing keywords")
    if result["missing_keywords"]:
        st.write(" ".join(f"`{k}`" for k in result["missing_keywords"]))
    else:
        st.write("No major keyword gaps found.")

    st.divider()
    st.subheader("Recommended improvements")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    if not result["improvements"]:
        st.write("No improvements returned.")
    for item in result["improvements"]:
        title = f"{icons[item['priority']]} {item['priority']} priority - {item['section']}"
        with st.expander(title, expanded=item["priority"] == "High"):
            st.markdown(f"**Issue:** {item['issue']}")
            st.markdown(f"**Fix:** {item['suggestion']}")
            if item["example"]:
                st.markdown("**Example:**")
                st.code(item["example"], language=None)

    report = json.dumps(result, indent=2)
    st.download_button(
        "Download report (JSON)",
        data=report,
        file_name="ats_report.json",
        mime="application/json",
    )


# --------------------------------------------------------------------------
# App
# --------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and tips to improve it.")

    with st.sidebar:
        st.header("Settings")
        sidebar_key = st.text_input(
            "Gemini API key",
            type="password",
            help="Optional if the key is already set in Streamlit secrets.",
        )
        st.markdown("[Get a free API key](https://aistudio.google.com/apikey)")
        st.markdown("---")
        st.caption(
            "Your resume is sent to Google's Gemini API for analysis and is not "
            "stored by this app."
        )

    api_key = get_api_key(sidebar_key)

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=160,
        placeholder="Paste the job description here for a tailored keyword match...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar or in Streamlit secrets.")
            st.stop()

        try:
            with st.spinner("Reading your resume..."):
                resume_text = extract_resume_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            st.stop()

        if len(resume_text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned "
                "image, ATS systems cannot read it either - export a text-based "
                "PDF or DOCX and try again."
            )
            st.stop()

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(api_key, resume_text, job_description)
        except json.JSONDecodeError:
            st.error("The model returned an unreadable response. Please try again.")
            st.stop()
        except Exception as exc:
            msg = str(exc)
            if "API key" in msg or "API_KEY" in msg or "403" in msg or "401" in msg:
                st.error("Your Gemini API key looks invalid. Please check it and try again.")
            elif "429" in msg or "quota" in msg.lower():
                st.error("Rate limit reached. Please wait a minute and try again.")
            else:
                st.error(f"Analysis failed: {msg}")
            st.stop()

        st.session_state["result"] = result
        st.session_state["resume_text"] = resume_text

    if "result" in st.session_state:
        render_results(st.session_state["result"])
        with st.expander("View extracted resume text"):
            st.text(st.session_state.get("resume_text", "")[:5000])


if __name__ == "__main__":
    main()
